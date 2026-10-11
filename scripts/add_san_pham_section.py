#!/usr/bin/env python3
"""
Insert PHẦN 6. SẢN PHẨM section into the existing KHKT report docx.
Uses python-docx high-level API for reliable formatting, then moves
newly created paragraphs/tables to BEFORE TÀI LIỆU THAM KHẢO.
"""
import sys
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INPUT_PATH = "/home/z/my-project/upload/Bao_cao_KHKT_Hoang_Bao_Lam_14_trang (3).docx"
OUTPUT_PATH = "/home/z/my-project/download/Bao_cao_KHKT_V11_with_San_pham.docx"

doc = Document(INPUT_PATH)

# Find target — TÀI LIỆU THAM KHẢO Heading 1
target_para = None
for p in doc.paragraphs:
    if p.style and p.style.name == 'Heading 1' and 'TÀI LIỆU THAM KHẢO' in p.text.upper():
        target_para = p
        break

if target_para is None:
    print("ERROR: TÀI LIỆU THAM KHẢO Heading 1 not found", file=sys.stderr)
    sys.exit(1)

print(f"Found target: {target_para.text!r} (style={target_para.style.name})")

# Reference existing styles to discover their actual IDs (python-docx handles this internally)
# Get the body element
body = doc.element.body

# ====== Helper: add a paragraph at end, then move it before target ======
new_elements_in_order = []

def add_para(text, style_name='Normal', bold=False, italic=False, alignment=None,
             font_name='Times New Roman', font_size=12, indent_first=True):
    """Create paragraph using python-docx API at end of doc, then return its element."""
    p = doc.add_paragraph(style=style_name)
    if alignment is not None:
        p.alignment = alignment
    # Line spacing 1.3x
    pf = p.paragraph_format
    pf.line_spacing = 1.3
    # First line indent for body (skip for headings)
    if indent_first and style_name == 'Normal':
        pf.first_line_indent = Cm(1.0)
    # Run with text
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    run.italic = italic
    # Set East Asian font
    rPr = run._element.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        run._element.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)

    # Detach and queue for insertion
    p_elem = p._element
    p_elem.getparent().remove(p_elem)
    new_elements_in_order.append(p_elem)
    return p

def add_page_break():
    """Add a paragraph with page break."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.3
    run = p.add_run()
    run.add_break(docx.enum.text.WD_BREAK.PAGE) if False else None
    # Use proper page break via XML
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._element.append(br)
    p_elem = p._element
    p_elem.getparent().remove(p_elem)
    new_elements_in_order.append(p_elem)
    return p

import docx.enum.text

def add_table(rows_data, headers):
    """Create a styled table at end of doc, then detach for insertion."""
    n_cols = len(headers)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        p.style = doc.styles['Normal']
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        rPr = run._element.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            run._element.insert(0, rPr)
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        # Shading
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'D9E2F3')
        tcPr.append(shd)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data rows
    for row_data in rows_data:
        cells = table.add_row().cells
        for i, val in enumerate(row_data):
            cell = cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            p.style = doc.styles['Normal']
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i < 2 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            rPr = run._element.find(qn('w:rPr'))
            if rPr is None:
                rPr = OxmlElement('w:rPr')
                run._element.insert(0, rPr)
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.insert(0, rFonts)
            rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            if i == 2:
                run.bold = True
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Set column widths
    for row in table.rows:
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(8.5)
        row.cells[2].width = Cm(3.0)

    # Detach for insertion
    t_elem = table._element
    t_elem.getparent().remove(t_elem)
    new_elements_in_order.append(t_elem)
    return table

# ====== Build new content ======

# Page break to start on new page
add_page_break()

# Heading 1 — PHẦN 6. SẢN PHẨM
add_para('PHẦN 6. SẢN PHẨM', style_name='Heading 1', bold=True, indent_first=False)

# Intro paragraph
add_para(
    'Bên cạnh báo cáo nghiên cứu, đề tài xây dựng một sản phẩm phần mềm minh hoạ trực quan: '
    'hệ thống dự báo cháy rừng vận hành theo thời gian thực, kết nối mô hình XGBoost v8.6 đã huấn luyện '
    'với dữ liệu thời tiết trực tiếp từ Open-Meteo API. Sản phẩm cho phép người dùng quan sát nguy cơ cháy '
    'trên toàn lãnh thổ Việt Nam, tương tác với từng ô lưới 0,5° và theo dõi lịch sử dự báo.'
)

# 6.1
add_para('6.1. Kiến trúc hệ thống và luồng dữ liệu', style_name='Heading 2', bold=True, indent_first=False)
add_para(
    'Hệ thống được xây dựng theo mô hình ứng dụng web Next.js 16 với API route chạy cùng tiến trình. '
    'Khi người dùng nhấp vào một ô lưới trên bản đồ Leaflet, frontend gọi đồng thời hai endpoint: '
    '(1) /api/weather để lấy thời tiết thực từ Open-Meteo dựa trên tọa độ ô đó, và (2) /api/predict '
    'để gửi toàn bộ 126 ô lưới Việt Nam qua chế độ batch cho mô hình XGBoost. API predict sẽ khởi tạo '
    'một subprocess Python (predict_v86.py), nạp mô hình XGBoost từ tệp xgb_modis_iid.json cùng bộ hiệu '
    'chuẩn isotonic và imputer, sau đó trả xác suất đã hiệu chuẩn cho từng ô. Toàn bộ luồng dữ liệu '
    'thực hiện trong khoảng ba giây cho 126 ô.'
)

# 6.2
add_para('6.2. Tính năng chính', style_name='Heading 2', bold=True, indent_first=False)
features = [
    ('Bản đồ tương tác 70/30', 'Bản đồ Leaflet chiếm 70% diện tích, với lớp ô lưới 0,5° phủ trên lãnh thổ Việt Nam, lớp điểm cháy MODIS/VIIRS với hiệu ứng phát sáng, và lớp thành phố chính. Bảng điều khiển 30% bên phải chứa các công cụ điều khiển.'),
    ('Chọn khoảng dự báo', 'Hỗ trợ ba khoảng dự báo t+1 (ngày mai), t+3 (sau ba ngày) và t+7 (sau bảy ngày). Ngày dự báo cụ thể được tính và hiển thị tự động (ví dụ: 11/10/2026).'),
    ('Thời tiết thực từ Open-Meteo', 'Khi người dùng nhấp vào ô lưới bất kỳ, hệ thống tự động gọi Open-Meteo API để lấy nhiệt độ, độ ẩm, tốc độ gió và lượng mưa hiện tại cho toạ độ đó, không cần thao tác thủ công.'),
    ('Tự động cập nhật', 'Trang tự làm mới dữ liệu thời tiết và chạy lại mô hình mỗi mười phút khi tab đang mở, có đồng hồ đếm ngược hiển thị thời gian đến lần cập nhật tiếp theo. Tự động tạm dừng khi tab ẩn để tiết kiệm băng thông.'),
    ('Lịch sử dự báo', 'Sidebar bên phải lưu mười lần dự báo gần nhất, ghi lại dấu thời gian, khoảng dự báo, điều kiện thời tiết đầu vào và xác suất trung bình toàn quốc.'),
    ('Trạng thái mô hình', 'Hệ thống hiển thị phiên bản mô hình (v8.6), AUROC IID (0,9463), độ trễ inference (khoảng 2.500–3.000 mili giây cho 126 ô) và nguồn dữ liệu thời tiết (Open-Meteo · ngày).'),
]
for title, desc in features:
    add_para(f'{title}. {desc}')

# 6.3
add_para('6.3. Công nghệ triển khai', style_name='Heading 2', bold=True, indent_first=False)
add_para(
    'Frontend được viết bằng TypeScript với khung React 19 và Next.js 16, sử dụng Tailwind CSS cho giao diện '
    'và thư viện react-leaflet cho bản đồ. Backend API route chạy trên Node.js, gọi subprocess Python 3 '
    'với thư viện xgboost và scikit-learn để nạp mô hình. Mô hình XGBoost v8.6 bao gồm 35 đặc trưng '
    '(gồm biến trễ, trung bình trượt, biến dẫn xuất vật lý), được huấn luyện trên 598.143 ô-ngày với nhãn '
    'MODIS Collection 6, sử dụng hiệu chuẩn isotonic regression để giảm sai số hiệu chuẩn (ECE) từ 0,0549 '
    'xuống 0,0090. Toàn bộ mã nguồn, mô hình và dữ liệu được đóng gói để tái lập lại đầy đủ.'
)

# 6.4
add_para('6.4. Đường dẫn truy cập và hướng dẫn sử dụng', style_name='Heading 2', bold=True, indent_first=False)
add_para(
    'Sản phẩm được triển khai dưới dạng ứng dụng web có thể truy cập qua trình duyệt. Người dùng truy cập '
    'trang giới thiệu, nhấp nút “Thử ngay” để vào giao diện bản đồ. Trên bản đồ, các ô màu thể hiện '
    'nguy cơ cháy từ thấp (xanh) đến rất cao (đỏ). Người dùng nhấp vào ô bất kỳ để xem thông tin chi tiết: '
    'toạ độ, vùng, xác suất cháy và mức nguy cơ. Có thể chuyển đổi khoảng dự báo t+1/t+3/t+7 bằng dropdown, '
    'thay đổi lớp bản đồ nền (OpenStreetMap, Esri vệ tinh, Esri địa hình), bật tắt lớp hiển thị '
    '(ô lưới, điểm cháy, thành phố). Hệ thống tự động cập nhật mỗi mười phút; người dùng cũng có thể '
    'nhấn nút “Cập nhật ngay” để làm mới theo yêu cầu.'
)

# 6.5 — table
add_para('6.5. Bảng tổng hợp tính năng', style_name='Heading 2', bold=True, indent_first=False)

table_rows = [
    ('Bản đồ tương tác', 'Leaflet + OpenStreetMap/Esri basemap, 126 ô lưới 0,5°, 190 điểm cháy', 'Đã triển khai'),
    ('Khoảng dự báo', 't+1, t+3, t+7 (chọn bằng dropdown, tính ngày tự động)', 'Đã triển khai'),
    ('Thời tiết thực', 'Open-Meteo API, tự lấy khi click ô (không cần thao tác thủ công)', 'Đã triển khai'),
    ('Tự động cập nhật', 'Mỗi 10 phút khi tab mở, tạm dừng khi tab ẩn', 'Đã triển khai'),
    ('Lịch sử dự báo', '10 lần gần nhất với timestamp và input thời tiết', 'Đã triển khai'),
    ('Mô hình thật', 'XGBoost v8.6, 35 đặc trưng, isotonic calibration, AUROC 0,9463', 'Đã nối thành công'),
    ('Cảnh báo OOD', 'Hiển thị warning khi chọn vùng Nam (AUROC 0,584)', 'Đã triển khai'),
    ('Giới hạn lãnh thổ', 'Polygon biên giới VN, loại bỏ cháy ngoài biển và nước láng giềng', 'Đã triển khai'),
]
add_table(table_rows, ['Nhóm tính năng', 'Chi tiết', 'Trạng thái'])

# ====== Now insert all new_elements_in_order BEFORE target_para ======
target_element = target_para._element
for elem in new_elements_in_order:
    target_element.addprevious(elem)

# Save
os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
doc.save(OUTPUT_PATH)
print(f"\n[OK] Saved: {OUTPUT_PATH}")
print(f"     Size: {os.path.getsize(OUTPUT_PATH):,} bytes")

# Verify
print("\n=== VERIFICATION ===")
doc2 = Document(OUTPUT_PATH)
print(f"Total paragraphs: {len(doc2.paragraphs)} (was 164, now +17)")
print(f"Total tables: {len(doc2.tables)} (was 6, now +1)")

print("\n--- PHẦN 6 / SẢN PHẨM related paragraphs ---")
for i, p in enumerate(doc2.paragraphs):
    if 'PHẦN 6' in p.text or 'SẢN PHẨM' in p.text.upper() or p.text.startswith('6.'):
        print(f"  [{i}] style={p.style.name!r} text={p.text[:80]!r}")

print("\n--- Surrounding TÀI LIỆU THAM KHẢO ---")
found_tl = False
for i, p in enumerate(doc2.paragraphs):
    if 'TÀI LIỆU THAM KHẢO' in p.text and p.style.name == 'Heading 1':
        found_tl = True
        print(f"  [{i}] style={p.style.name!r} text={p.text!r}")
        # Show paragraphs before
        print("  --- 5 paragraphs BEFORE ---")
        for j in range(max(0, i-5), i):
            print(f"    [{j}] style={doc2.paragraphs[j].style.name!r} text={doc2.paragraphs[j].text[:60]!r}")
        break

# Also verify the new table is in the right place (before TÀI LIỆU THAM KHẢO)
print("\n--- Tables ---")
for i, t in enumerate(doc2.tables):
    rows = len(t.rows)
    cols = len(t.columns) if rows > 0 else 0
    first_cell = t.rows[0].cells[0].text if rows > 0 else ''
    print(f"  Table {i}: {rows} rows × {cols} cols, first cell: {first_cell[:50]!r}")
